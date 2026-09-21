import hashlib,json
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
B=Path(__file__).resolve().parent
R=Path('C:/personal_webpage/my-portfolio/APR_code_gen/Part2_Create_test')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
data=json.loads((B/'results.json').read_text())
# Remove inherited historical per-test statuses/timing. Current raw logs are authoritative.
for r in data:
 for k in ('subtests','seconds','batch'):r.pop(k,None)
(B/'results.json').write_text(json.dumps(data,indent=2))
old=json.loads((R/'test_validation/consolidated_runs.json').read_text())
changes=json.loads((B/'changes.json').read_text())
names={c[0] for c in changes}
assert {(r['case'],r['version'],r['variant']) for r in data}=={(r['case'],r['version'],r['variant']) for r in old if r['case'] in names}
assert len(data)==240
for r in data:
 assert r['source_sha256']==sha(Path(r['source']))
 assert r['test_sha256']==sha(Path(r['test']))
 assert Path(r['log']).is_file()
 assert r['status']=='PASS' if r['variant']=='fixed' else r['status']!='PASS'
manifest=[]
for name in sorted(names):
 for filename in ['buggy.py','fixed.py','test.py','original_question.txt']:
  a=R/'reconstructed_cases'/name/filename;b=B/'cases'/name/filename
  manifest.append(dict(case=name,file=filename,before=sha(a),after=sha(b)))
  if filename in ['buggy.py','original_question.txt']:assert sha(a)==sha(b)
assert sum(m['before']!=m['after'] for m in manifest)==7
path=B/'outputs/repair01/Quantum_APR_REVIEW_FIXED_Analysis.xlsx'
if path.exists():
 original=load_workbook(R/'test_validation/Quantum_APR_REVIEW_FIXED_Analysis.xlsx')
 w=load_workbook(path)
 for s in original:
  for row in s:
   for c in row:
    if s.title=='Case analysis' and c.coordinate in ['A5','A6']:continue
    assert c.value==w[s.title][c.coordinate].value,(s.title,c.coordinate)
 assert w['Repair runs'].max_row==241
 cached=load_workbook(path,data_only=True)
 assert cached['Repair results']['B4'].value==120
 assert cached['Repair results']['D4'].value==0
manifest.append(dict(file='Quantum_APR_REVIEW_FIXED_Analysis.xlsx',before=sha(R/'test_validation/Quantum_APR_REVIEW_FIXED_Analysis.xlsx'),after=sha(path) if path.exists() else None))
(B/'manifest.json').write_text(json.dumps(manifest,indent=2))
print(Counter((r['variant'],r['status']) for r in data))
print('Coverage, hashes, preservation, and report QA passed.')
