"""Read-only source coverage, preservation, and saved-XLSX audit."""
from copy import copy
import json
from pathlib import Path
import re
import openpyxl
BASE=Path(__file__).resolve().parent
OUT=BASE/'outputs/01a07c80-batch06'
old=openpyxl.load_workbook(OUT/'input_snapshot.xlsx')
new=openpyxl.load_workbook(OUT/'Quantum_APR_Validation.xlsx')
values=openpyxl.load_workbook(OUT/'Quantum_APR_Validation.xlsx',data_only=True)
assert old.sheetnames==new.sheetnames
for name in old.sheetnames:
    assert old[name].freeze_panes==new[name].freeze_panes
    assert str(old[name].merged_cells)==str(new[name].merged_cells)
    assert old[name].sheet_view.showGridLines==new[name].sheet_view.showGridLines
    assert list(old[name].tables)==list(new[name].tables)
    for col in old[name].column_dimensions:
        assert old[name].column_dimensions[col].width==new[name].column_dimensions[col].width
for name,rows,cols in [('Summary',57,10),('Version results',710,16),('Fixed failures',86,9),('Environments',151,11)]:
    for row in range(1,rows+1):
        for col in range(1,cols+1):
            a=old[name].cell(row,col);b=new[name].cell(row,col)
            changed=name=='Summary' and row in (3,4) and col==1
            if not changed: assert a.value==b.value,(name,a.coordinate)
            assert copy(a.font)==copy(b.font),(name,a.coordinate,'font')
            assert a.number_format==b.number_format,(name,a.coordinate,'number format')
for row in range(60,90):
    for col in range(1,11):
        assert old['Summary'].cell(row,col).value==new['Summary'].cell(row+16,col).value
for row in range(161,166):
    for col in range(1,12):
        assert old['Environments'].cell(row,col).value==new['Environments'].cell(row+27,col).value
for sheet in new:
    for row in sheet:
        for cell in row: assert cell.data_type!='e',(sheet.title,cell.coordinate,cell.value)
for row in range(7,74):
    assert not re.search('[\uac00-\ud7a3]',str(new['Summary'].cell(row,10).value))
assert [values['Summary'].cell(74,c).value for c in range(3,7)]==[937,802,135,0]
for name,table,ref in [('Version results','AllVersionResults','A1:P938'),('Fixed failures','AllFixedFailures','A1:I136'),('Environments','AllEnvironments','A1:K178')]:
    assert new[name].tables[table].ref==ref
def read(name): return json.loads((BASE/name).read_text(encoding='utf-8'))
selection=read('selection.json');meta=read('metadata.json');pairs=read('new_pairs.json');data=read('consolidated_data.json')
source=openpyxl.load_workbook(selection['workbook'],data_only=True)
valid={}
for sheet in source:
    for row in sheet.values:
        if str(row[1]).strip().lower()=='p':
            number=int(re.search(r'\d+',str(row[0])).group())
            valid.setdefault(number,set()).update(v.strip() for v in row[3].split(','))
actual={int(re.search(r'\d+',c['case']).group()):c for c in data['cases']}
assert set(valid)==set(actual), ('Remaining or extra IDs',set(valid)^set(actual))
assert len(actual)==len(data['cases'])==67
root=Path(selection['cases'][0]['source_dir']).parent
for number,c in actual.items():
    if (root/('issue_%03d_se'%number)).exists(): assert c['case'].endswith('_se'),c
    assert set(c['versions'])==valid[number],('Version coverage',c['case'])
for index,c in enumerate(selection['cases'],58):
    assert str(source[c['sheet']].cell(c['row'],2).value).lower()=='p'
    assert [v.strip() for v in source[c['sheet']].cell(c['row'],4).value.split(',')]==c['versions']
    assert values['Summary'].cell(index,2).value==c['case']
    assert values['Summary'].cell(index,3).value==len(c['versions'])
    assert values['Summary'].cell(index,4).value==sum(p['fixed']=='PASS' for p in pairs if p['case']==c['case'])
    assert new['Summary'].cell(index,10).value==meta['j_notes'].get(c['case'])
for row,p in zip(values['Version results'].iter_rows(min_row=711,values_only=True),pairs):
    assert row[:9]==(6,p['case'],p['version'],p['buggy'],p['fixed'],p['category'],p['checks_passed'],p['checks_failed'],p['checks_errors'])
    assert row[12]==p['bug_reason'] and (row[13] or '')==p['fix_reason']
print('EXPORT AUDIT PASSED: all 67 valid case IDs and versions covered; no remaining cases. Historical values, formulas, J notes, fonts, widths and panes preserved.')
