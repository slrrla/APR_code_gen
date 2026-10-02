"""Prepare exact-version jobs and preserve the original workbook container."""
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
IDS = {772, 790, 927, 858, 901, 921}
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def main():
    original = BASE / 'validity_before.xlsx'
    authored = BASE / 'validity_updated.xlsx'
    before = openpyxl.load_workbook(original)
    after = openpyxl.load_workbook(authored)
    changes = []
    cases = []
    for row in before.active:
        if row[0].value not in IDS:
            continue
        number = int(row[0].value)
        cases.append({'case': 'issue_%03d' % number,
                      'versions': [v.strip() for v in row[3].value.split(',')],
                      'validity_before': row[1].value, 'validity_after': 'p',
                      'workbook_cell': row[1].coordinate})
        assert after.active[row[1].coordinate].value == 'p'
        changes.append(row[1].coordinate)
    assert len(cases) == 6
    unexpected = [(a.coordinate, a.value, after.active[a.coordinate].value)
                  for row in before.active for a in row
                  if a.value != after.active[a.coordinate].value
                  and a.coordinate not in changes]
    assert not unexpected, unexpected

    # Artifact Tool authored the requested values, but its XLSX serializer
    # normalizes unrelated original styles. Transfer only the approved value
    # delta into the original OOXML container to preserve all other entries.
    # Reuse its existing shared 'p' string; preserve styles and XML lexically.
    with ZipFile(original) as source:
        strings = ET.fromstring(source.read('xl/sharedStrings.xml'))
        values = [''.join(si.itertext()) for si in strings.findall('s:si', NS)]
        p_index = values.index('p')
        sheet_name = 'xl/worksheets/sheet1.xml'
        xml = source.read(sheet_name).decode('utf-8')
        for cell in changes:
            pattern = r'(<c\b[^>]*\br="' + cell + r'"[^>]*>)(.*?)(</c>)'
            match = re.search(pattern, xml, flags=re.DOTALL)
            assert match and 't="s"' in match.group(1), cell
            body, count = re.subn(r'<v>\d+</v>', '<v>%d</v>' % p_index,
                                 match.group(2), count=1)
            assert count == 1, cell
            xml = xml[:match.start(2)] + body + xml[match.end(2):]
        final = BASE / 'validity_ready.xlsx'
        with ZipFile(final, 'w') as destination:
            for entry in source.infolist():
                content = xml.encode('utf-8') if entry.filename == sheet_name else source.read(entry.filename)
                destination.writestr(entry, content)
    ready = openpyxl.load_workbook(final)
    assert before.sheetnames == ready.sheetnames
    for row in before.active:
        for old in row:
            new = ready.active[old.coordinate]
            assert new.value == ('p' if old.coordinate in changes else old.value), old.coordinate
            assert old._style == new._style, old.coordinate
    with ZipFile(original) as a, ZipFile(final) as b:
        assert a.namelist() == b.namelist()
        assert all(a.read(name) == b.read(name) for name in a.namelist() if name != sheet_name)
    selection = {'workbook': r'C:\Users\haha9\Downloads\Valid_Cases_104 (1).xlsx',
                 'workbook_before_sha256': hashlib.sha256(original.read_bytes()).hexdigest(),
                 'workbook_after_sha256': hashlib.sha256(final.read_bytes()).hexdigest(),
                 'cases': sorted(cases, key=lambda c: c['case'])}
    (BASE / 'selection.json').write_text(json.dumps(selection, indent=2), encoding='utf-8')
    source_hashes = {}
    for case in cases:
        number = int(case['case'].split('_')[1])
        for path in [ROOT / case['case'] / 'buggy.py', ROOT / case['case'] / 'fixed.py',
                     ROOT / 'failpass_fixes' / ('case_%d_new_fix.py' % number)]:
            if path.exists():
                source_hashes[str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
    (BASE / 'source_hashes_before.json').write_text(json.dumps(source_hashes, indent=2), encoding='utf-8')
    print('Verified six P changes; all other original workbook ZIP entries and cell styles preserved.')
    print('Prepared %d exact case/version pairs.' % sum(len(c['versions']) for c in cases))


if __name__ == '__main__':
    main()
