"""Merge the Artifact Tool-authored one-cell change into the original XLSX."""
import hashlib
import json
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

BASE = Path(__file__).resolve().parent
selection = json.loads((BASE / 'selection.json').read_text(encoding='utf-8'))
item = selection['cases'][0]
assert len(selection['cases']) == 1 and item['case'] == 'issue_827'
original = BASE / 'validity_before.xlsx'
authored = BASE / 'validity_updated.xlsx'
before = openpyxl.load_workbook(original)
after = openpyxl.load_workbook(authored)
cell = item['workbook_cell']
assert before.active[cell].value == item['validity_before']
assert after.active[cell].value == 'p'
assert all(a.value == after.active[a.coordinate].value for row in before.active
           for a in row if a.coordinate != cell)

# Artifact Tool export normalizes unrelated original styles. Its approved
# value delta alone is projected into the original OOXML container; all other
# original ZIP entries, cell styles, dimensions and metadata stay byte-identical.
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
sheet_name = 'xl/worksheets/sheet1.xml'
final = BASE / 'validity_ready.xlsx'
with ZipFile(original) as source:
    root = ET.fromstring(source.read('xl/sharedStrings.xml'))
    strings = [''.join(node.itertext()) for node in root.findall('s:si', ns)]
    p_index = strings.index('p')
    xml = source.read(sheet_name).decode('utf-8')
    match = re.search(r'(<c\b[^>]*\br="' + cell + r'"[^>]*>)(.*?)(</c>)', xml, re.DOTALL)
    assert match and 't="s"' in match.group(1)
    body, count = re.subn(r'<v>\d+</v>', '<v>%d</v>' % p_index, match.group(2), count=1)
    assert count == 1
    xml = xml[:match.start(2)] + body + xml[match.end(2):]
    with ZipFile(final, 'w') as output:
        for entry in source.infolist():
            output.writestr(entry, xml.encode('utf-8') if entry.filename == sheet_name else source.read(entry.filename))

ready = openpyxl.load_workbook(final)
assert ready.sheetnames == before.sheetnames
for row in before.active:
    for old in row:
        new = ready.active[old.coordinate]
        assert new.value == ('p' if old.coordinate == cell else old.value)
        assert old._style == new._style
with ZipFile(original) as a, ZipFile(final) as b:
    assert a.namelist() == b.namelist()
    assert all(a.read(name) == b.read(name) for name in a.namelist() if name != sheet_name)
selection['workbook_after_sha256'] = hashlib.sha256(final.read_bytes()).hexdigest()
selection['p_count_after'] = sum(str(row[1].value).strip().lower() == 'p' for row in ready.active)
(BASE / 'selection.json').write_text(json.dumps(selection, indent=2), encoding='utf-8')
print('Verified only %s changed to p; original formatting and all other ZIP entries preserved.' % cell)
print('P count: %d -> %d.' % (selection['p_count_before'], selection['p_count_after']))
