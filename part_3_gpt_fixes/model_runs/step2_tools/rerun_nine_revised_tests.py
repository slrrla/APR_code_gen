"""User-run targeted rerun of nine revised APR oracles and workbook merge.

This script makes no model API requests and never edits tests or generated fixes.
"""

import argparse
import collections
import copy
import json
import os
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill


CASES = (
    "issue_096", "issue_453", "issue_497", "issue_622", "issue_769",
    "issue_876", "issue_877", "issue_958", "issue_985",
)
MODELS = (
    "gpt-5.6-luna", "gpt-6-luna", "gpt-6.1-sol", "gpt-6-astra",
    "claude-opus-5-5", "claude-sonnet-5-5", "claude-haiku-5-5",
)
BOOK_NAMES = {
    "gpt-5.6-luna": "73_cases_5.6_luna_fixes (5).xlsx",
    "gpt-6-luna": "73_cases_6_luna_fixes (2).xlsx",
    "gpt-6.1-sol": "73_cases_6.1_sol_fixes (2).xlsx",
    "gpt-6-astra": "73_cases_6_astra_fixes (1).xlsx",
    "claude-opus-5-5": "73_cases_claude_opus_5_5_reviewed.xlsx",
    "claude-sonnet-5-5": "73_cases_claude_sonnet_5_5_reviewed.xlsx",
    "claude-haiku-5-5": "73_cases_claude_haiku_5_5_reviewed.xlsx",
}
DEFAULT_ROOT = (Path.home() / "Downloads" / "APR_code_gen-main (3)" /
                "APR_code_gen-main" / "Part2_Create_test" / "model_runs")
BATCH_NAME = "revised_nine_2026_10_20"


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def tools_module(root):
    sys.path.insert(0, str(root / "step2_tools"))
    import run_tests  # noqa: E402
    original_child_environment = run_tests.child_environment

    def child_environment(py):
        env = original_child_environment(py)
        for name in (
            "ANTHROPIC_API_KEY", "IBM_QUANTUM_TOKEN", "QISKIT_IBM_TOKEN",
            "QISKIT_IBM_RUNTIME_TOKEN", "QISKIT_IBM_RUNTIME_API_TOKEN",
        ):
            env.pop(name, None)
        return env

    run_tests.child_environment = child_environment
    return run_tests


def row_keys(sheet):
    found = {}
    for row in sheet.iter_rows(min_row=2):
        if len(row) < 2:
            continue
        case, version = row[0].value, row[1].value
        if not (isinstance(case, str) and case.startswith("issue_")):
            continue
        key = (case, str(version))
        if key in found:
            raise ValueError(f"Duplicate workbook row: {key}")
        found[key] = row[0].row
    return found


def plan_for(root, model, workbook, tools):
    manifest = read_json(root / "_shared" / "manifest.json")
    entries = [item for item in manifest["cases"] if item["case"] in CASES]
    if {item["case"] for item in entries} != set(CASES):
        raise ValueError("The manifest does not contain exactly the nine requested cases")
    valid_path = root / "_shared" / "Valid_Cases_104.xlsx"
    valid_sheet = openpyxl.load_workbook(valid_path, read_only=True, data_only=True).worksheets[0]
    plan = []
    for item in entries:
        case = item["case"]
        listed = re.findall(r"\d+\.\d+\.\d+", str(valid_sheet.cell(item["valid_row"], 4).value))
        if listed != item["versions"]:
            raise ValueError(f"Valid_Cases_104 version mismatch: {case}")
        label = str(valid_sheet.cell(item["valid_row"], 1).value)
        if int(re.search(r"\d+", label)[0]) != int(case.split("_")[1]):
            raise ValueError(f"Valid_Cases_104 case mismatch: {case}")
        case_dir = Path(item["source_case_dir"])
        test = case_dir / "test.py"
        human = case_dir / "fixed.py"
        name = case.split("_")[1]
        fix_name = f"luna_fix_{name}.py" if model == "gpt-5.6-luna" else f"fix_{name}.py"
        model_fix = root / model / "generation" / case / fix_name
        for path in (test, human, model_fix):
            if not path.is_file():
                raise FileNotFoundError(path)
        status_path = model_fix.parent / "status.json"
        if status_path.is_file():
            status = read_json(status_path)
            if (status.get("status") != "CREATED" or status.get("model") != model or
                    status.get("fix_sha256") != tools.digest(model_fix)):
                raise ValueError(f"Generated fix provenance changed: {model_fix}")
        plan.append({
            "case": case, "versions": listed, "test": str(test),
            "test_sha256": tools.digest(test),
            "targets": {"HUMAN_FIXED": str(human), "MODEL_FIXED": str(model_fix)},
            "target_sha256": {"HUMAN_FIXED": tools.digest(human),
                              "MODEL_FIXED": tools.digest(model_fix)},
        })
    if sum(len(item["versions"]) for item in plan) != 150:
        raise ValueError("Expected exactly 150 assigned version pairs")
    if not workbook.is_file():
        raise FileNotFoundError(f"Workbook missing: {workbook}")
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    if book.sheetnames != ["Exit results", "Detailed reasons", "Summary"]:
        raise ValueError(f"Unexpected workbook sheets: {book.sheetnames}")
    expected = {(item["case"], version) for item in plan for version in item["versions"]}
    for sheet_name in ("Exit results", "Detailed reasons"):
        keys = row_keys(book[sheet_name])
        if len(keys) != 1015 or {key for key in keys if key[0] in CASES} != expected:
            raise ValueError(f"Workbook {sheet_name} does not have the expected 1015 rows/assignments")
    book.close()
    helpers = {str(path): tools.digest(path) for path in
               sorted((root.parent / "reconstructed_cases").glob("*oracle*.py"))}
    return {"model": model, "cases": plan, "helpers": helpers, "source_valid": str(valid_path),
            "workbook": str(workbook), "timeout_seconds": 900}


def result_path(batch, entry, version, target):
    return batch / "logs" / entry["case"] / version / (target + ".json")


def collect(plan, batch, tools):
    records = []
    for entry in plan["cases"]:
        for version in entry["versions"]:
            for target in ("HUMAN_FIXED", "MODEL_FIXED"):
                path = result_path(batch, entry, version, target)
                if not path.is_file():
                    continue
                record = read_json(path)
                if (record["case"], record["version"], record["target"]) != (entry["case"], version, target):
                    raise ValueError(f"Wrong result key: {path}")
                if (record["test_sha256"] != entry["test_sha256"] or
                        record["target_sha256"] != entry["target_sha256"][target] or
                        tools.digest(Path(record["log"])) != record["log_sha256"]):
                    raise ValueError(f"Result input/log hash mismatch: {path}")
                records.append(record)
    return records


def check_plan(batch, plan, tools):
    path = batch / "plan.json"
    if path.is_file():
        if read_json(path) != plan:
            raise ValueError("A test, fix, workbook path or assigned version changed; use a new batch")
    else:
        if batch.exists() and any(batch.iterdir()):
            raise ValueError(f"Batch contains files but no plan: {batch}")
        batch.mkdir(parents=True, exist_ok=True)
        tools.write_json(path, plan)


def run(plan, batch, tools, env_root, max_runs, resume_interrupted):
    snapshots = tools.probe_environments(plan, env_root)
    check_plan(batch, plan, tools)
    lock = batch / "runner.lock"
    with lock.open("x", encoding="utf-8") as handle:
        json.dump({"pid": os.getpid(), "started": datetime.now(timezone.utc).isoformat()}, handle)
    try:
        done = {(r["case"], r["version"], r["target"]) for r in collect(plan, batch, tools)}
        attempted = 0
        for entry in plan["cases"]:
            for version in entry["versions"]:
                for target in ("HUMAN_FIXED", "MODEL_FIXED"):
                    key = (entry["case"], version, target)
                    if key in done:
                        continue
                    if max_runs is not None and attempted >= max_runs:
                        print(f"Stopped at --max-runs {max_runs}; run the same command again to resume.")
                        return
                    marker = batch / "logs" / entry["case"] / version / (target + ".inflight.json")
                    if marker.exists():
                        if not resume_interrupted:
                            raise RuntimeError(f"Interrupted marker: {marker}. Verify no child is running, then use --resume-interrupted")
                        old_log = marker.with_name(target + ".txt")
                        if old_log.is_file():
                            old_log.rename(old_log.with_name(old_log.stem + ".interrupted_" +
                                                             datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + ".txt"))
                        marker.unlink()
                    print(f"[{len(done) + 1}/300] {entry['case']} / {version} / {target}", flush=True)
                    record = tools.execute_one(entry, version, target, snapshots[version], batch, plan["timeout_seconds"])
                    done.add(key)
                    attempted += 1
                    print(f"{record['classification']}; exit={record['exit_code']}; {record['seconds']}s", flush=True)
    finally:
        lock.unlink(missing_ok=True)
        records = collect(plan, batch, tools)
        tools.write_json(batch / "results.json", {"plan": plan, "records": records})
        print(f"Saved {len(records)}/300 executions in {batch}")


def status_palette(sheet, column):
    palette = {}
    for row in sheet.iter_rows(min_row=2):
        cell = row[column - 1]
        if cell.value in ("PASS", "ERROR", "FAIL") and cell.value not in palette:
            palette[cell.value] = copy.copy(cell.fill)
    return palette


def set_status(cell, value, palette):
    cell.value = value
    if value in palette:
        cell.fill = copy.copy(palette[value])


def merge(plan, batch, workbook, tools):
    records = collect(plan, batch, tools)
    if len(records) != 300:
        raise RuntimeError(f"Only {len(records)}/300 executions finished; workbook unchanged")
    by_key = {(r["case"], r["version"], r["target"]): r for r in records}
    book = openpyxl.load_workbook(workbook)
    raw = book["Exit results"]
    detail = book["Detailed reasons"]
    summary = book["Summary"]
    raw_rows = row_keys(raw)
    detail_rows = row_keys(detail)
    old_unaffected = {
        sheet.title: {(case, version): tuple(sheet.cell(row, col).value for col in range(3, 9))
                      for (case, version), row in row_keys(sheet).items() if case not in CASES}
        for sheet in (raw, detail)
    }
    palettes = {(raw.title, col): status_palette(raw, col) for col in (4, 5)}
    palettes.update({(detail.title, col): status_palette(detail, col) for col in (4, 6)})
    if detail.cell(1, 8).value in (None, ""):
        detail.cell(1, 8).value = "Test Sensitivity / Review"
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    for entry in plan["cases"]:
        for version in entry["versions"]:
            key = (entry["case"], version)
            rr, dr = raw_rows[key], detail_rows[key]
            human = by_key[key + ("HUMAN_FIXED",)]
            model = by_key[key + ("MODEL_FIXED",)]
            set_status(raw.cell(rr, 4), human["raw_result"], palettes[(raw.title, 4)])
            set_status(raw.cell(rr, 5), model["raw_result"], palettes[(raw.title, 5)])
            raw.cell(rr, 6).value = human["exit_code"]
            raw.cell(rr, 7).value = model["exit_code"]
            set_status(detail.cell(dr, 4), human["classification"], palettes[(detail.title, 4)])
            set_status(detail.cell(dr, 6), model["classification"], palettes[(detail.title, 6)])
            detail.cell(dr, 5).value = "Revised test.py: " + human["reason"]
            detail.cell(dr, 7).value = "Revised test.py: " + model["reason"]
            note = detail.cell(dr, 8)
            note.value = (None if model["classification"] == "PASS" else
                          "REVIEW: revised test result; compare question, fix and log before calling this test-sensitive.")
            note.fill = review_fill if note.value else PatternFill(fill_type=None)
    counts = collections.Counter()
    human_counts = collections.Counter()
    buggy_counts = collections.Counter()
    all_detail = row_keys(detail)
    for key, row in all_detail.items():
        counts[(key[0], detail.cell(row, 6).value)] += 1
        human_counts[detail.cell(row, 4).value] += 1
        buggy_counts[raw.cell(raw_rows[key], 3).value] += 1
    total = collections.Counter({status: sum(counts[(case, status)] for case in {k[0] for k in all_detail})
                                 for status in ("PASS", "ERROR", "FAIL")})
    seen_cases = set()
    total_row = None
    metrics_section = str(summary.cell(1, 7).value).strip().upper() == "METRIC"
    for row in range(2, summary.max_row + 1):
        case = summary.cell(row, 1).value
        if case in {k[0] for k in all_detail}:
            seen_cases.add(case)
            summary.cell(row, 2).value = sum(counts[(case, status)] for status in ("PASS", "ERROR", "FAIL"))
            for col, status in ((3, "PASS"), (4, "ERROR"), (5, "FAIL")):
                summary.cell(row, col).value = counts[(case, status)]
        elif case == "TOTAL":
            total_row = row
            summary.cell(row, 2).value = len(all_detail)
            for col, status in ((3, "PASS"), (4, "ERROR"), (5, "FAIL")):
                summary.cell(row, col).value = total[status]
        label = summary.cell(row, 7).value
        if not isinstance(label, str):
            continue
        normalized = re.sub(r"\s+", " ", label.replace(":", " ")).strip().upper()
        if normalized == "METRIC":
            metrics_section = True
            continue
        if not metrics_section:
            continue
        if normalized == "ASSIGNED VERSION ROWS":
            summary.cell(row, 8).value = len(all_detail)
        elif normalized.startswith("BUGGY "):
            summary.cell(row, 8).value = buggy_counts[normalized.removeprefix("BUGGY ")]
        elif normalized.startswith("HUMAN "):
            summary.cell(row, 8).value = human_counts[normalized.removeprefix("HUMAN ")]
        elif normalized in ("PASS", "ERROR", "FAIL") or normalized.startswith(("LUNA ", "SOL ", "ASTRA ", "CLAUDE ")):
            summary.cell(row, 8).value = total[normalized.split()[-1]]
    if len(seen_cases) != 73 or total_row is None or sum(total.values()) != 1015:
        raise ValueError("Summary did not reconcile to 73 cases / 1015 versions")
    # Prior side tables describe the old oracle. Retain their layout but withdraw stale flags.
    case_numbers = {int(case.split("_")[1]) for case in CASES}
    for row in detail.iter_rows(min_col=10):
        for cell in row:
            if cell.value in case_numbers and isinstance(cell.value, (int, float)):
                adjacent = detail.cell(cell.row, cell.column + 1)
                if isinstance(adjacent.value, (int, float)):
                    adjacent.value = 0
                    detail.cell(cell.row, cell.column + 2).value = "Revised test: test sensitivity pending manual review."
    for sheet in (raw, detail):
        after = {(case, version): tuple(sheet.cell(row, col).value for col in range(3, 9))
                 for (case, version), row in row_keys(sheet).items() if case not in CASES}
        if after != old_unaffected[sheet.title]:
            raise ValueError(f"Unrelated case rows changed in {sheet.title}")
    backup_dir = batch / "workbook_backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / (workbook.stem + "_before_revised_tests_" +
                           datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + ".xlsx")
    shutil.copy2(workbook, backup)
    temporary = workbook.with_name(workbook.stem + ".revised_tmp.xlsx")
    try:
        book.save(temporary)
        verify = openpyxl.load_workbook(temporary, read_only=True, data_only=True)
        assert verify.sheetnames == ["Exit results", "Detailed reasons", "Summary"]
        check_raw = {(row[0], str(row[1])): row for row in verify["Exit results"].values
                     if len(row) >= 7 and isinstance(row[0], str) and row[0].startswith("issue_")}
        check_detail = {(row[0], str(row[1])): row for row in verify["Detailed reasons"].values
                        if len(row) >= 7 and isinstance(row[0], str) and row[0].startswith("issue_")}
        assert len(check_raw) == len(check_detail) == 1015
        for entry in plan["cases"]:
            for version in entry["versions"]:
                key = (entry["case"], version)
                human = by_key[key + ("HUMAN_FIXED",)]
                model = by_key[key + ("MODEL_FIXED",)]
                assert check_raw[key][3:7] == (human["raw_result"], model["raw_result"],
                                              human["exit_code"], model["exit_code"])
                assert check_detail[key][3:7] == (human["classification"], "Revised test.py: " + human["reason"],
                                                 model["classification"], "Revised test.py: " + model["reason"])
        assert sum(int(verify["Summary"].cell(total_row, col).value or 0) for col in (3, 4, 5)) == 1015
        verify.close()
        book.close()
        temporary.replace(workbook)
    finally:
        temporary.unlink(missing_ok=True)
    print(f"Updated: {workbook}")
    print(f"Backup:  {backup}")
    print(f"Model totals: {dict(total)}")
    print("Changed-case non-passes are flagged for manual semantic review in column H.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model", required=True, choices=MODELS)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--env-root", type=Path, default=Path.home() / "miniconda3" / "envs")
    parser.add_argument("--max-runs", type=int)
    parser.add_argument("--resume-interrupted", action="store_true")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--plan", action="store_true")
    mode.add_argument("--preflight", action="store_true")
    mode.add_argument("--run", action="store_true")
    mode.add_argument("--merge", action="store_true")
    args = parser.parse_args()
    if args.max_runs is not None and args.max_runs < 1:
        parser.error("--max-runs must be positive")
    workbook = args.workbook or (args.root / args.model / BOOK_NAMES[args.model])
    tools = tools_module(args.root)
    plan = plan_for(args.root, args.model, workbook, tools)
    batch = args.root / args.model / "testing" / BATCH_NAME
    if args.plan:
        for entry in plan["cases"]:
            print(entry["case"], ", ".join(entry["versions"]), Path(entry["test"]).name)
        print(f"9 cases, 150 assigned pairs, 300 human/model executions; workbook: {workbook}")
        return
    if args.preflight:
        tools.probe_environments(plan, args.env_root)
        print("PREFLIGHT OK; no case tests run and no workbook modified")
        return
    if args.run:
        run(plan, batch, tools, args.env_root, args.max_runs, args.resume_interrupted)
        return
    if not (batch / "plan.json").is_file() or read_json(batch / "plan.json") != plan:
        raise ValueError("Revision plan missing or inputs changed; workbook unchanged")
    merge(plan, batch, workbook, tools)


if __name__ == "__main__":
    main()
