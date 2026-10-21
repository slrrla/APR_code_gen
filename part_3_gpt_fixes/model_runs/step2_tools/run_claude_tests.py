"""Run one Claude repair model against APR (3) tests and assigned versions.

Copy this file into Part2_Create_test/model_runs/step2_tools. This runner uses
the existing run_tests.py there for environment probing, isolation, logging,
and execution classification. It never calls a model API.
"""

import argparse
import importlib.util
import os
import re
import shutil
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


HERE = Path(__file__).resolve().parent
ROOT = Path(os.environ["APR_MODEL_RUNS_ROOT"]).resolve() if "APR_MODEL_RUNS_ROOT" in os.environ else HERE.parent
VALID = Path.home() / "Downloads" / "Valid_Cases_104 (2).xlsx"
BASELINE = Path.home() / "Downloads" / "73_cases_5.6_luna_fixes (3).xlsx"
MODELS = ("claude-opus-5-5", "claude-sonnet-5-5", "claude-haiku-5-5")
TARGETS = ("HUMAN_FIXED", "CLAUDE_FIXED")
TEST_PREFERENCE = ("new_test.py", "new-test.py", "test_new.py", "test-new.py")

spec = importlib.util.spec_from_file_location("apr_test_tools", ROOT / "step2_tools" / "run_tests.py")
if spec is None or spec.loader is None:
    raise RuntimeError("Cannot load the existing step2_tools/run_tests.py")
tools = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tools)

original_child_environment = tools.child_environment


def child_environment(py):
    env = original_child_environment(py)
    for name in (
        "ANTHROPIC_API_KEY", "IBM_QUANTUM_TOKEN", "QISKIT_IBM_TOKEN",
        "QISKIT_IBM_RUNTIME_TOKEN", "QISKIT_IBM_RUNTIME_API_TOKEN",
    ):
        env.pop(name, None)
    return env


tools.child_environment = child_environment


def make_plan(model):
    manifest = tools.read_json(ROOT / "_shared" / "manifest.json")
    entries = sorted(manifest["cases"], key=lambda item: int(re.search(r"\d+", item["case"]).group()))
    if len(entries) != 73:
        raise RuntimeError(f"Expected 73 cases, found {len(entries)}")

    valid = load_workbook(VALID, data_only=True, read_only=True).worksheets[0]
    baseline = load_workbook(BASELINE, data_only=True, read_only=True).worksheets[0]
    buggy = {}
    for row in baseline.iter_rows(min_row=2, values_only=True):
        if row and isinstance(row[0], str) and row[0].startswith("issue_"):
            key = (row[0], str(row[1]))
            if key in buggy:
                raise RuntimeError(f"Duplicate buggy baseline row: {key}")
            buggy[key] = row[2]

    cases = []
    case_root = (ROOT.parent / "reconstructed_cases").resolve()
    for entry in entries:
        case = entry["case"]
        versions = re.findall(r"\d+\.\d+\.\d+", str(valid.cell(entry["valid_row"], 4).value))
        row_label = str(valid.cell(entry["valid_row"], 1).value)
        if int(re.search(r"\d+", row_label).group()) != int(re.search(r"\d+", case).group()):
            raise RuntimeError(f"Valid_104 row points to a different case: {case}")
        if versions != entry["versions"] or len(versions) != len(set(versions)):
            raise RuntimeError(f"Valid_104 versions changed: {case}")
        if any(version not in tools.ENVIRONMENTS for version in versions):
            raise RuntimeError(f"No Qiskit environment mapping for {case}")

        folder = Path(entry["source_case_dir"]).resolve()
        if not folder.is_relative_to(case_root) or folder.name != case:
            raise RuntimeError(f"Case is not inside APR (3): {folder}")
        updated = [folder / name for name in TEST_PREFERENCE if (folder / name).is_file()]
        if len(updated) > 1:
            raise RuntimeError(f"More than one updated test exists for {case}: {updated}")
        test = updated[0] if updated else folder / "test.py"
        human = folder / "fixed.py"
        generation = ROOT / model / "generation" / case
        status = tools.read_json(generation / "status.json")
        fix = generation / status["fix_file"]
        for path in (test, human, fix):
            if not path.is_file():
                raise FileNotFoundError(path)
        if status["status"] != "CREATED" or status["model"] != model:
            raise RuntimeError(f"Claude generation is incomplete or belongs to another model: {case}")
        if tools.digest(fix) != status["fix_sha256"]:
            raise RuntimeError(f"Claude fix changed since generation: {case}")
        if status["input_sha256"] != entry["input_sha256"]["messages.json"]:
            raise RuntimeError(f"Claude prompt input mismatch: {case}")
        recorded = {}
        for version in versions:
            if (case, version) not in buggy:
                raise RuntimeError(f"Missing buggy baseline: {case}/{version}")
            recorded[version] = buggy[(case, version)]
        cases.append({
            "case": case, "versions": versions, "test": str(test),
            "test_sha256": tools.digest(test),
            "targets": {"HUMAN_FIXED": str(human), "CLAUDE_FIXED": str(fix)},
            "target_sha256": {"HUMAN_FIXED": tools.digest(human), "CLAUDE_FIXED": tools.digest(fix)},
            "buggy_recorded": recorded,
        })

    if sum(len(item["versions"]) for item in cases) != 1015:
        raise RuntimeError("Assigned case/version total is not 1015")
    helpers = {str(path): tools.digest(path) for path in sorted(case_root.glob("*oracle*.py"))}
    return {
        "model": model, "cases": cases, "helpers": helpers,
        "source_valid": str(VALID), "source_baseline": str(BASELINE),
        "buggy_policy": "Copied from the 5.6 baseline; not rerun.",
        "classification": "Log-based PASS/ERROR/FAIL; semantic and test-sensitivity review still required.",
    }


def collect(plan, batch):
    records = []
    for entry in plan["cases"]:
        for version in entry["versions"]:
            for target in TARGETS:
                record_path = batch / "logs" / entry["case"] / version / (target + ".json")
                if not record_path.is_file():
                    continue
                record = tools.read_json(record_path)
                if (record["case"], record["version"], record["target"]) != (entry["case"], version, target):
                    raise RuntimeError(f"Wrong result key: {record_path}")
                if record["test_sha256"] != entry["test_sha256"]:
                    raise RuntimeError(f"Test changed after run: {record_path}")
                if record["target_sha256"] != entry["target_sha256"][target]:
                    raise RuntimeError(f"Target changed after run: {record_path}")
                if tools.digest(Path(record["log"])) != record["log_sha256"]:
                    raise RuntimeError(f"Log changed after run: {record_path}")
                records.append(record)
    return records


def safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_workbook(plan, batch, records):
    lookup = {(row["case"], row["version"], row["target"]): row for row in records}
    if len(lookup) != len(records):
        raise RuntimeError("Duplicate test records")
    wb = Workbook()
    exits = wb.active
    exits.title = "Exit results"
    detail = wb.create_sheet("Detailed reasons")
    summary = wb.create_sheet("Summary")
    exits.append(("Case Number", "Qiskit Version", "Buggy Result", "Human Fix Result", "Claude Fix Result", "Human Exit Code", "Claude Exit Code"))
    detail.append(("Case Number", "Qiskit Version", "Buggy Result", "Human Fix Result", "Human Reason", "Claude Fix Result", "Claude Reason", "Test Sensitivity Review", "Test File", "Human Log", "Claude Log"))
    summary.append(("Case Number", "Versions Tested", "Claude PASS", "Claude ERROR", "Claude FAIL", "Claude NOT TESTED"))
    overall = Counter()
    for entry in plan["cases"]:
        counts = Counter()
        for version in entry["versions"]:
            human = lookup.get((entry["case"], version, "HUMAN_FIXED"), {})
            claude = lookup.get((entry["case"], version, "CLAUDE_FIXED"), {})
            hc = human.get("classification", "NOT TESTED")
            cc = claude.get("classification", "NOT TESTED")
            buggy = entry["buggy_recorded"][version]
            exits.append((entry["case"], version, buggy, human.get("raw_result", "NOT TESTED"),
                          claude.get("raw_result", "NOT TESTED"), human.get("exit_code"), claude.get("exit_code")))
            detail.append(tuple(map(safe, (entry["case"], version, buggy, hc, human.get("reason", "Pending"),
                                           cc, claude.get("reason", "Pending"), "", Path(entry["test"]).name,
                                           human.get("log", ""), claude.get("log", "")))))
            counts[cc] += 1
            overall[("HUMAN_FIXED", hc)] += 1
            overall[("CLAUDE_FIXED", cc)] += 1
        summary.append((entry["case"], sum(counts[key] for key in ("PASS", "ERROR", "FAIL")),
                        counts["PASS"], counts["ERROR"], counts["FAIL"], counts["NOT TESTED"]))
    summary.append(("TOTAL", sum(summary.cell(r, 2).value for r in range(2, summary.max_row + 1)),
                    *(sum(summary.cell(r, c).value for r in range(2, summary.max_row + 1)) for c in range(3, 7))))
    summary["H1"], summary["I1"] = "Metric", "Count"
    for i, (target, category) in enumerate(((t, c) for t in TARGETS for c in ("PASS", "ERROR", "FAIL", "NOT TESTED")), 2):
        summary.cell(i, 8, f"{target} {category}")
        summary.cell(i, 9, overall[(target, category)])
    summary["H12"], summary["I12"] = "Assigned pairs", 1015
    summary["H13"], summary["I13"] = "Completed processes", len(records)
    summary["H15"], summary["I15"] = "Classification scope", "Test/log results only; review flagged failures against the question."
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{ws.cell(1, ws.max_column).column_letter}{ws.max_row}"
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="D9E5F2")
            cell.font = Font(bold=True)
        for column, width in {"A": 20, "B": 18, "C": 18, "D": 20, "E": 65, "F": 20, "G": 90, "H": 62, "I": 28, "J": 58, "K": 58}.items():
            ws.column_dimensions[column].width = width
    for row in detail.iter_rows(min_row=2):
        for cell in (row[4], row[6], row[7]):
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = batch / ("73_cases_" + plan["model"].replace("-", "_").replace(".", "_") + "_fixes.xlsx")
    temp = output.with_suffix(".tmp.xlsx")
    wb.save(temp)
    temp.replace(output)
    print(f"Workbook saved: {output}")
    print(f"Verified: {sum(len(entry['versions']) for entry in plan['cases'])} assigned rows, {len(records)} logged executions.")
    print(dict(overall))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model", choices=MODELS, required=True)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--plan", action="store_true")
    mode.add_argument("--preflight", action="store_true")
    mode.add_argument("--run", action="store_true")
    mode.add_argument("--report-only", action="store_true")
    parser.add_argument("--max-runs", type=int)
    parser.add_argument("--timeout", type=int, default=900)
    parser.add_argument("--resume-interrupted", action="store_true")
    parser.add_argument("--env-root", type=Path, default=Path.home() / "miniconda3" / "envs")
    args = parser.parse_args()
    if args.timeout < 1 or (args.max_runs is not None and args.max_runs < 1):
        parser.error("--timeout and --max-runs must be positive")
    plan = make_plan(args.model)
    batch = ROOT / args.model / "testing" / "all73"
    if args.plan:
        for entry in plan["cases"]:
            print(entry["case"], ", ".join(entry["versions"]), Path(entry["test"]).name)
        print("73 cases; 1015 assigned pairs; 2030 Human/Claude test processes.")
        return
    if args.report_only:
        if tools.read_json(batch / "plan.json") != plan:
            raise RuntimeError("Saved plan differs from current inputs; do not merge evidence")
        records = collect(plan, batch)
        tools.save_results(plan, batch, records)
        export_workbook(plan, batch, records)
        return
    snapshots = tools.probe_environments(plan, args.env_root)
    if args.preflight:
        print("PREFLIGHT OK; no case tests ran.")
        return
    batch.mkdir(parents=True, exist_ok=True)
    lock = batch / "runner.lock"
    with lock.open("x", encoding="utf-8") as handle:
        handle.write(f"{os.getpid()} {datetime.now(timezone.utc).isoformat()}\n")
    try:
        plan_path = batch / "plan.json"
        env_path = batch / "environments.json"
        if plan_path.exists():
            if tools.read_json(plan_path) != plan or tools.read_json(env_path) != snapshots:
                raise RuntimeError("Plan or environments changed; preserve this batch and investigate")
        else:
            tools.write_json(plan_path, plan)
            tools.write_json(env_path, snapshots)
        records = collect(plan, batch)
        completed = {(r["case"], r["version"], r["target"]) for r in records}
        attempted = 0
        for entry in plan["cases"]:
            for version in entry["versions"]:
                for target in TARGETS:
                    key = (entry["case"], version, target)
                    if key in completed:
                        continue
                    if args.max_runs is not None and attempted >= args.max_runs:
                        break
                    folder = batch / "logs" / entry["case"] / version
                    marker = folder / (target + ".inflight.json")
                    if marker.exists():
                        if not args.resume_interrupted:
                            raise RuntimeError(f"Interrupted run marker: {marker}; check old process, then use --resume-interrupted")
                        old_log = folder / (target + ".txt")
                        if old_log.exists():
                            old_log.rename(folder / (target + f".interrupted_{time.time_ns()}.txt"))
                    print(f"[{len(records) + 1}/2030] {entry['case']} / {version} / {target}", flush=True)
                    result = tools.execute_one(entry, version, target, snapshots[version], batch, args.timeout)
                    records.append(result)
                    completed.add(key)
                    attempted += 1
                    print(f"{result['classification']}; exit={result['exit_code']}; {result['seconds']} seconds", flush=True)
                    tools.save_results(plan, batch, records)
                if args.max_runs is not None and attempted >= args.max_runs:
                    break
            if args.max_runs is not None and attempted >= args.max_runs:
                break
        records = collect(plan, batch)
        tools.save_results(plan, batch, records)
        export_workbook(plan, batch, records)
        print(f"Saved {len(records)}/2030 executions to {batch}")
    except KeyboardInterrupt:
        records = collect(plan, batch)
        tools.save_results(plan, batch, records)
        export_workbook(plan, batch, records)
        print("Interrupted. Completed records preserved; resume after checking any inflight marker.")
    finally:
        lock.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
