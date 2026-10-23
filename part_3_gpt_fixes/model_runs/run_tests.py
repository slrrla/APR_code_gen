"""Step 2: assigned versions only, checkpointed tests, and three-sheet reports."""
import argparse
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import time
import uuid
from datetime import datetime, timezone

ENVIRONMENTS = {
    '0.25.0': 'qiskit_0250', '0.25.1': 'qiskit_0251_exact',
    '0.25.2': 'qiskit_0252_exact', '0.25.3': 'qiskit_0253_exact',
    '0.45.0': 'qiskit_045', '0.45.1': 'qiskit_0451_exact',
    '0.45.2': 'qiskit_0452_exact', '0.45.3': 'qiskit_0453_exact',
    '0.46.0': 'qiskit_0460_exact', '0.46.1': 'qiskit_0461_exact',
    '0.46.2': 'qiskit_0462', '0.46.3': 'qiskit_0463',
    '1.0.0': 'qiskit_100', '1.0.1': 'qiskit_101_exact', '1.0.2': 'qiskit_102_exact',
    '1.1.0': 'qiskit_110_exact', '1.1.1': 'qiskit_111', '1.1.2': 'qiskit_112_exact',
    '1.2.0': 'qiskit_120_exact', '1.2.1': 'qiskit_121_exact',
    '1.2.2': 'qiskit_122_exact', '1.2.4': 'qiskit_124',
    '2.0.0': 'qiskit_200_exact', '2.2.0': 'qiskit_220_exact',
    '2.3.0': 'qiskit_230_exact', '2.4.0': 'qiskit_240_exact', '2.5.0': 'qiskit_250',
}
TARGETS = ('BUGGY', 'HUMAN_FIXED', 'LUNA_FIXED')
HERE = Path(__file__).resolve().parent


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def atomic(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(text, encoding='utf-8')
    temp.replace(path)


def write(path, data):
    atomic(path, json.dumps(data, indent=2, ensure_ascii=False) + '\n')


def fingerprint(data):
    return hashlib.sha256(json.dumps(data, sort_keys=True).encode()).hexdigest()


def now():
    return datetime.now(timezone.utc).isoformat()


def norm(path):
    return os.path.normcase(os.path.abspath(path))


def make_plan(root, model, batch):
    import openpyxl
    manifest = read(root / '_shared' / 'manifest.json')
    assert model == 'gpt-6-luna', 'This runner is scoped to gpt-6-luna first 35.'
    selected = sorted(manifest['cases'], key=lambda e: int(re.search(r'\d+', e['case'])[0]))[:35]
    assert len(selected) == 35 and selected[-1]['case'] == 'issue_596'
    valid_path = Path(manifest['source_valid_workbook'])
    valid = openpyxl.load_workbook(valid_path, data_only=True).worksheets[0]
    entries = []
    for original in selected:
        case = original['case']
        versions = re.findall(r'\d+\.\d+\.\d+', str(valid.cell(original['valid_row'], 4).value))
        label = valid.cell(original['valid_row'], 1).value
        assert int(re.search(r'\d+', str(label))[0]) == int(re.search(r'\d+', case)[0])
        assert versions == original['versions'], f'valid_104 changed: {case}; review assignments first.'
        assert len(versions) == len(set(versions)) and set(versions) <= set(ENVIRONMENTS)
        folder = Path(original['source_case_dir'])
        tests = [folder / name for name in ('new_test.py', 'new-test.py', 'test.py') if (folder / name).is_file()]
        if not tests:
            raise FileNotFoundError(f'No test for {case}')
        test = tests[0]
        assert 'MUT' in test.read_text(encoding='utf-8-sig'), f'Test does not expose MUT: {test}'
        paths = {'BUGGY': folder / 'buggy.py', 'HUMAN_FIXED': folder / 'fixed.py',
                 'LUNA_FIXED': root / model / 'generation' / case / f'fix_{case.removeprefix("issue_")}.py'}
        for p in paths.values():
            assert p.is_file(), f'Missing target: {p}'
        status = read(paths['LUNA_FIXED'].parent / 'status.json')
        assert status['model'] == model and status['fix_sha256'] == sha(paths['LUNA_FIXED'])
        helpers = {str(p): sha(p) for p in folder.parent.glob('*oracle*.py')}
        entry = {'case': case, 'versions': versions, 'valid_row': original['valid_row'],
                 'test': str(test), 'test_sha256': sha(test), 'helpers': helpers,
                 'targets': {k: {'path': str(p), 'sha256': sha(p)} for k, p in paths.items()},
                 'timeout_seconds': 1800 if case == 'issue_225' else 600}
        entries.append(entry)
    plan = {'model': model, 'cases': entries, 'batch': batch, 'targets': list(TARGETS),
            'valid_workbook': str(valid_path), 'valid_workbook_sha256': sha(valid_path),
            'environment_names': ENVIRONMENTS,
            'harness_sha256': sha(HERE / 'test_harness.py'),
            'runner_sha256': sha(Path(__file__)),
            'classification': 'unittest evidence; non-PASS reasons need review before attributing model correctness'}
    assert sum(len(e['versions']) for e in entries) == 501
    return plan


def save_plan(out, plan):
    path = out / 'plan.json'
    if path.exists():
        assert read(path) == plan, 'Source/configuration changed; use a new --batch name to keep previous results.'
    else:
        write(path, plan)
        # Preserve test/target evidence without running or rewriting any source.
        for entry in plan['cases']:
            snapshots = {entry['test']: 'test_used.py'}
            snapshots.update({v['path']: k.lower() + '.py' for k, v in entry['targets'].items()})
            for src, name in snapshots.items():
                dest = out / 'source_snapshots' / entry['case'] / name
                dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(src, dest)


def child_environment(python, working, case_root):
    env = os.environ.copy()
    for key in list(env):
        upper = key.upper()
        if upper.startswith(('OPENAI_', 'OPEN_AI_', 'AZURE_OPENAI_')) or upper in ('PYTHONHOME', 'PYTHONPATH'):
            env.pop(key, None)
    envroot = python.parent
    env.update({'PYTHONIOENCODING': 'utf-8', 'PYTHONUTF8': '1', 'PYTHONDONTWRITEBYTECODE': '1',
                'PYTHONNOUSERSITE': '1', 'PYTHONPATH': str(case_root),
                'CONDA_PREFIX': str(envroot), 'CONDA_DEFAULT_ENV': envroot.name,
                'PATH': os.pathsep.join(str(envroot / p) for p in ('', 'Scripts', 'Library/bin')) + os.pathsep + env.get('PATH', '')})
    home = working / 'home'
    temp = working / 'temp'
    home.mkdir(parents=True, exist_ok=True)
    temp.mkdir(parents=True, exist_ok=True)
    env.update({'HOME': str(home), 'USERPROFILE': str(home), 'IPYTHONDIR': str(home / '.ipython'),
                'TEMP': str(temp), 'TMP': str(temp), 'MPLCONFIGDIR': str(home / '.matplotlib')})
    return env


def kill_tree(process):
    if os.name == 'nt':
        subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'], stdout=subprocess.DEVNULL,
                       stderr=subprocess.DEVNULL, timeout=30)
    else:
        process.kill()
    if process.poll() is None:
        process.kill()
    process.wait(timeout=30)


def probe_environments(plan, out):
    names = sorted({v for c in plan['cases'] for v in c['versions']}, key=lambda x: tuple(map(int, x.split('.'))))
    result = {}
    for version in names:
        python = Path.home() / 'miniconda3' / 'envs' / ENVIRONMENTS[version] / 'python.exe'
        if not python.is_file():
            raise FileNotFoundError(f'Missing interpreter: {python}')
        work = out / 'environment_checks' / version
        env = child_environment(python, work, Path(plan['cases'][0]['test']).parent.parent)
        code = (
            'import json,sys; from importlib.metadata import version,PackageNotFoundError\n'
            'def get(name):\n'
            ' try: return version(name)\n'
            ' except PackageNotFoundError: return None\n'
            'import qiskit\n'
            'print("APR_ENV="+json.dumps({"python":sys.version,"qiskit":get("qiskit"),'
            '"aer":get("qiskit-aer"),"numpy":get("numpy"),"ipython":get("ipython"),'
            '"runtime":get("qiskit-ibm-runtime"),"python_constraint":get("python-constraint")}))'
        )
        p = subprocess.run([str(python), '-B', '-c', code], env=env, cwd=work,
                           capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=60)
        atomic(work / 'probe.log', p.stdout + '\n' + p.stderr)
        lines = [s for s in p.stdout.splitlines() if s.startswith('APR_ENV=')]
        if p.returncode or len(lines) != 1:
            raise RuntimeError(f'Environment probe failed: {version}. See {work / "probe.log"}')
        info = json.loads(lines[0].split('=', 1)[1])
        if info['qiskit'] != version:
            raise RuntimeError(f'Version mismatch: {ENVIRONMENTS[version]} has {info["qiskit"]}, expected {version}. No tests started.')
        info.update(executable=str(python), environment=ENVIRONMENTS[version])
        result[version] = info
        print(f'Environment {version}: exact Qiskit confirmed; Aer={info["aer"]}', flush=True)
    previous = out / 'environments.json'
    if previous.exists() and read(previous) != result:
        raise RuntimeError('Environment versions changed since this batch. Use a new --batch name; do not mix results.')
    write(previous, result)
    return result


def classify(code, evidence, target, timed_out=False):
    """Keep runtime stops distinct from completed-output/interface assertions."""
    if timed_out:
        return 'ERROR', '(runtime error: test process exceeded the recorded time limit; timeout is not proof of wrong output.)', True
    if code is not None and (code < 0 or code >= 0x80000000):
        return 'ERROR', f'(runtime error: native process terminated with Windows status 0x{code & 0xffffffff:08X}.)', True
    if not evidence or not evidence.get('finished') or not evidence.get('tests_run'):
        return 'ERROR', '(runtime error: test suite did not complete a recorded test; inspect the console log.)', True
    events = evidence.get('events', [])
    problems = []
    for event in events:
        kind = event.get('kind')
        if kind == 'success':
            continue
        message = re.sub(r'\s+', ' ', event.get('message', '')).strip()
        detail = f'{event.get("type", kind)}: {message}'[:650]
        frames = event.get('frames', [])
        inside_target = any(norm(f['file']) == norm(target) for f in frames)
        native = 'SIMULATOR_PROCESS_ERROR:' in message
        assertion = event.get('type') == 'AssertionError'
        last = frames[-1]['file'] if frames else ''
        offline = any(s in message.lower() for s in ('offline', 'network access', 'online account'))
        if kind in ('skip', 'expected_failure', 'unexpected_success'):
            problems.append(('ERROR', f'Test coverage is incomplete or exceptional ({kind}): {message}'))
        elif native or event.get('type') == 'TimeoutExpired':
            problems.append(('ERROR', detail))
        elif kind == 'harness_error':
            problems.append(('ERROR', 'Test setup failed: ' + detail))
        elif inside_target and (not assertion or norm(last) == norm(target) or offline):
            problems.append(('ERROR', detail))
        elif kind == 'failure' or event.get('type') in ('KeyError', 'AttributeError', 'TypeError', 'ValueError', 'SyntaxError'):
            problems.append(('FAIL', 'Test output/interface check failed: ' + detail))
        else:
            problems.append(('ERROR', 'Test/helper execution stopped; attribution requires review: ' + detail))
    if any(p[0] == 'ERROR' for p in problems):
        reason = next(p[1] for p in problems if p[0] == 'ERROR')
        if any(p[0] == 'FAIL' for p in problems):
            reason += '; output/assertion checks also failed.'
        return 'ERROR', '(runtime error: ' + reason + ')', True
    if problems:
        return 'FAIL', problems[0][1], True
    if code != 0:
        return 'ERROR', f'(runtime error: exit code {code} despite no captured failing test; inspect log.)', True
    return 'PASS', f'The assigned test completed successfully ({evidence["tests_run"]} tests).', False


def run_one(plan, entry, version, target, envinfo, out):
    job_folder = out / 'logs' / entry['case'] / version / target
    folder = job_folder / ('attempt_' + uuid.uuid4().hex[:12])
    folder.mkdir(parents=True, exist_ok=True)
    python = Path(envinfo['executable'])
    source = Path(entry['targets'][target]['path'])
    test = Path(entry['test'])
    for path, expected in [(source, entry['targets'][target]['sha256']), (test, entry['test_sha256'])]:
        assert sha(path) == expected, f'Input changed: {path}'
    for path, expected in entry['helpers'].items():
        assert sha(Path(path)) == expected, f'Helper changed: {path}'
    work = folder / 'working'
    env = child_environment(python, work, test.parent.parent)
    env['MUT'] = str(source)
    evidence_path = folder / 'unittest.json'
    command = [str(python), '-B', '-u', str(HERE / 'test_harness.py'), str(test), str(evidence_path)]
    started = time.monotonic()
    timeout = False
    code = None
    launch_error = None
    log = folder / 'console.log'
    with log.open('w', encoding='utf-8') as handle:
        handle.write(f'{entry["case"]} / {version} / {target}\nTest: {test}\nTarget: {source}\n')
        handle.flush()
        process = None
        try:
            process = subprocess.Popen(command, cwd=work, env=env, stdout=handle, stderr=subprocess.STDOUT,
                                       creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            try:
                code = process.wait(timeout=entry['timeout_seconds'])
            except subprocess.TimeoutExpired:
                timeout = True
                kill_tree(process)
                code = process.returncode
                handle.write('\nRUNNER_TIMEOUT\n')
        except KeyboardInterrupt:
            if process is not None and process.poll() is None:
                kill_tree(process)
            handle.write('\nRUNNER_INTERRUPTED: no completed record; this run will be retried on resume.\n')
            raise
        except OSError as exc:
            launch_error = str(exc)
            handle.write('\nRUNNER_LAUNCH_ERROR: ' + launch_error)
    evidence = read(evidence_path) if evidence_path.is_file() else None
    outcome, reason, review = classify(code, evidence, str(source), timeout)
    if launch_error:
        reason = '(runtime error: could not launch test interpreter: ' + launch_error + ')'
    assert sha(source) == entry['targets'][target]['sha256'], 'Target file changed during execution; stop and review.'
    assert sha(test) == entry['test_sha256'], 'Test changed during execution; stop and review.'
    result = {'case': entry['case'], 'version': version, 'target': target, 'model': plan['model'],
              'exit_code': code, 'raw_result': 'PASS' if code == 0 else 'FAIL',
              'classification': outcome, 'reason': reason, 'review_required': review,
              'duration_seconds': round(time.monotonic() - started, 3), 'completed_utc': now(),
              'log': str(log), 'evidence': str(evidence_path), 'test': str(test), 'source': str(source),
              'test_sha256': entry['test_sha256'], 'source_sha256': entry['targets'][target]['sha256'],
              'timeout_seconds': entry['timeout_seconds'], 'timed_out': timeout,
              'plan_sha256': fingerprint(plan), 'environment': envinfo}
    write(job_folder / 'result.json', result)
    return result


def collect(plan, out):
    records = []
    for entry in plan['cases']:
        for version in entry['versions']:
            for target in TARGETS:
                path = out / 'logs' / entry['case'] / version / target / 'result.json'
                if path.is_file():
                    row = read(path)
                    assert (row['case'], row['version'], row['target']) == (entry['case'], version, target)
                    assert row['plan_sha256'] == fingerprint(plan), 'Cannot combine different batch plans.'
                    records.append(row)
    return records


def reports(plan, out, excel=False):
    records = collect(plan, out)
    payload = {'plan': plan, 'records': records}
    write(out / 'results.json', payload)
    fields = ['case', 'version', 'target', 'model', 'exit_code', 'raw_result', 'classification',
              'reason', 'review_required', 'duration_seconds', 'completed_utc', 'test', 'source', 'log']
    buffer = io.StringIO(newline='')
    writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction='ignore')
    writer.writeheader()
    for row in records:
        writer.writerow({k: ("'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v)
                         for k, v in row.items()})
    try:
        atomic(out / 'results.csv', '\ufeff' + buffer.getvalue())
    except PermissionError:
        print('Close results.csv in Excel. Individual result.json files remain saved.', flush=True)
    if excel:
        node = os.environ.get('APR_NODE_EXE') or shutil.which('node')
        if not node:
            raise FileNotFoundError('Node.js missing: set APR_NODE_EXE or add node to PATH.')
        command = [str(node), str(HERE / 'export_test_results.mjs'), str(out / 'results.json'),
                   str(out / f'{plan["model"]}_{plan["batch"]}_results.xlsx')]
        p = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=180)
        atomic(out / 'workbook_export.log', p.stdout + '\n' + p.stderr)
        if p.returncode:
            print('Workbook export needs attention; logs/JSON/CSV are saved. See workbook_export.log.', flush=True)
        else:
            print(p.stdout.strip(), flush=True)
    return records


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--root', type=Path, default=HERE)
    ap.add_argument('--model', default='gpt-6-luna', choices=['gpt-6-luna'])
    ap.add_argument('--batch', default='first35')
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument('--plan-only', action='store_true')
    mode.add_argument('--preflight', action='store_true')
    mode.add_argument('--report-only', action='store_true')
    ap.add_argument('--max-runs', type=int, help='Stop after this many NEW test processes; resume skips finished runs')
    args = ap.parse_args()
    assert re.fullmatch(r'[A-Za-z0-9_-]+', args.batch), 'Invalid batch name'
    if args.max_runs is not None and args.max_runs < 1:
        ap.error('--max-runs must be positive')
    out = args.root / args.model / 'testing' / args.batch
    plan = make_plan(args.root, args.model, args.batch)
    save_plan(out, plan)
    print('35 cases / 501 assigned case-version pairs / 1503 BUGGY + HUMAN_FIXED + LUNA_FIXED tests.', flush=True)
    if args.plan_only or args.report_only:
        records = reports(plan, out, excel=True)
        print(f'Completed tests: {len(records)}/1503. No test processes launched.')
        return
    environments = probe_environments(plan, out)
    if args.preflight:
        print('Exact-version preflight passed. No case tests launched; runtime dependencies may still fail during tests.')
        return
    lock = out / 'testing.lock'
    try:
        with lock.open('x', encoding='utf-8') as handle:
            json.dump({'pid': os.getpid(), 'started_utc': now()}, handle)
    except FileExistsError:
        raise RuntimeError(f'Another test batch may be running. Inspect PID in {lock} before removing a stale lock.') from None
    finished = {(r['case'], r['version'], r['target']) for r in collect(plan, out)}
    count = 0
    try:
        for entry in plan['cases']:
            for version in entry['versions']:
                for target in TARGETS:
                    key = (entry['case'], version, target)
                    if key in finished:
                        continue
                    print(f'[{len(finished)+1}/1503] {entry["case"]} / {version} / {target}', flush=True)
                    result = run_one(plan, entry, version, target, environments[version], out)
                    finished.add(key)
                    count += 1
                    print(f'  {result["classification"]}; exit={result["exit_code"]}; {result["duration_seconds"]:.1f}s', flush=True)
                    reports(plan, out)
                    if args.max_runs and count >= args.max_runs:
                        return
            reports(plan, out, excel=True)
    except KeyboardInterrupt:
        print('\nStopped. Completed tests are retained; rerun the same command to resume.', flush=True)
    finally:
        lock.unlink(missing_ok=True)
        reports(plan, out, excel=True)
        print(f'Results folder: {out}', flush=True)


if __name__ == '__main__':
    main()
