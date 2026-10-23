"""Test only gpt-6.1-sol. Inherit the recorded buggy/human baseline."""
import argparse
import csv
import ctypes
import io
import json
import os
import re
import shutil
import subprocess
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import runner_core as core

MODEL = "gpt-6.1-sol"
ROOT = Path.home() / "Downloads/APR_code_gen-main (3)/APR_code_gen-main/Part2_Create_test/model_runs"
BASELINE = ROOT / "gpt-6-luna/gpt_6_luna_73_cases_revised_reasoning.xlsx"
VALID = Path.home() / "Downloads/Valid_Cases_104 (2).xlsx"
BOOK_NAME = "73_cases_6_1_sol_fixes.xlsx"
TARGET = "SOL_FIXED"


def case_rows(sheet):
    result = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not str(row[0] or "").startswith("issue_"):
            continue
        key = (row[0], str(row[1]))
        if key in result:
            raise ValueError(f"Duplicate baseline row: {key}")
        result[key] = row
    return result


def make_plan(args):
    manifest = core.read_json(args.root / "_shared/manifest.json")
    entries = sorted(manifest["cases"], key=lambda e: int(re.search(r"\d+", e["case"])[0]))
    if len(entries) != 73 or len({e["case"] for e in entries}) != 73:
        raise ValueError("Expected the existing 73-case generation manifest.")
    with args.baseline.open("rb") as handle:
        book = openpyxl.load_workbook(handle, read_only=True, data_only=True)
        old_exit = case_rows(book["Exit results"])
        old_detail = case_rows(book["Detailed reasons"])
        book.close()
    if set(old_exit) != set(old_detail):
        raise ValueError("Baseline sheets have different case/version rows.")
    valid = openpyxl.load_workbook(args.valid, read_only=True, data_only=True)
    valid_rows = list(valid.worksheets[0].iter_rows(values_only=True))
    valid.close()
    plan_cases = []
    expected_keys = set()
    cases_root = (args.root.parent / "reconstructed_cases").resolve()
    for entry in entries:
        case = entry["case"]
        source_row = valid_rows[entry["valid_row"] - 1]
        versions = re.findall(r"\d+\.\d+\.\d+", str(source_row[3]))
        if int(re.search(r"\d+", str(source_row[0]))[0]) != int(re.search(r"\d+", case)[0]):
            raise ValueError(f"Valid_104 row identity changed: {case}")
        if versions != entry["versions"] or len(versions) != len(set(versions)):
            raise ValueError(f"Valid_104 version assignment changed: {case}")
        if any(v not in core.ENVIRONMENTS for v in versions):
            raise ValueError(f"No environment mapping for {case}: {versions}")
        case_dir = cases_root / case
        if case_dir != Path(entry["source_case_dir"]).resolve():
            raise ValueError(f"Case must come from APR (3): {case}")
        alternatives = [case_dir / name for name in ("new_test.py", "new-test.py") if (case_dir / name).is_file()]
        if len(alternatives) > 1:
            raise ValueError(f"Two new-test alternatives exist: {case_dir}")
        test = alternatives[0] if alternatives else case_dir / "test.py"
        human = case_dir / "fixed.py"
        generation = args.root / MODEL / "generation" / case
        state = core.read_json(generation / "status.json")
        sol = generation / state["fix_file"]
        for path in (test, human, sol):
            if not path.is_file():
                raise FileNotFoundError(path)
        if sol.parent.resolve() != generation.resolve():
            raise ValueError(f"Unexpected generated target path: {sol}")
        if (state["model"], state["reasoning_effort"], state["status"]) != (MODEL, "medium", "CREATED"):
            raise ValueError(f"Wrong generation model/effort/status: {case}")
        if core.digest(sol) != state["fix_sha256"]:
            raise ValueError(f"Generated Sol fix changed: {case}")
        if state["input_sha256"] != entry["input_sha256"]["messages.json"]:
            raise ValueError(f"Generation input mismatch: {case}")
        test_hash, human_hash = core.digest(test), core.digest(human)
        inherited = {}
        for version in versions:
            key = (case, version)
            expected_keys.add(key)
            raw, detail = old_exit[key], old_detail[key]
            if raw[2] != detail[2] or detail[3] not in {"PASS", "ERROR", "FAIL"}:
                raise ValueError(f"Invalid recorded baseline status: {key}")
            if raw[3] != ("PASS" if raw[5] == 0 else "FAIL"):
                raise ValueError(f"Human exit-code mapping mismatch: {key}")
            if (raw[3] == "PASS") != (detail[3] == "PASS"):
                raise ValueError(f"Human raw/reviewed status mismatch: {key}")
            if detail[8] != test.name:
                raise ValueError(f"Test differs from carried human baseline: {key}; {detail[8]} -> {test.name}")
            log = Path(detail[10])
            if not log.is_file():
                raise FileNotFoundError(f"Missing carried human log: {log}")
            meta_file = log.with_suffix(".json")
            provenance = "Historical human metadata and log verified."
            if meta_file.is_file():
                meta = core.read_json(meta_file)
                if (meta["case"], meta["version"], meta["target"]) != (case, version, "HUMAN_FIXED"):
                    raise ValueError(f"Human metadata identity mismatch: {key}")
                if meta["classification"] != detail[3] or meta["exit_code"] != raw[5]:
                    raise ValueError(f"Human log/workbook result mismatch: {key}")
                if meta["test_sha256"] != test_hash or meta["target_sha256"] != human_hash:
                    raise ValueError(f"Human code or test changed since baseline: {key}")
                if core.digest(log) != meta["log_sha256"]:
                    raise ValueError(f"Human log changed: {key}")
            elif case == "issue_790":
                with (log.parent.parent / "results.csv").open(encoding="utf-8-sig", newline="") as stream:
                    matching = [r for r in csv.DictReader(stream) if (r["Case"], r["QiskitVersion"], r["Target"]) == (case, version, "HUMAN_FIXED")]
                if len(matching) != 1 or int(matching[0]["ExitCode"]) != raw[5] or matching[0]["TestResult"] != detail[3]:
                    raise ValueError(f"Case 790 revised-human rerun disagrees: {key}")
                provenance = "Revised-human rerun CSV/log verified; source hashes captured at Sol setup."
            else:
                raise FileNotFoundError(f"Missing human execution metadata: {meta_file}")
            inherited[version] = {
                "buggy_result": raw[2], "human_raw_result": raw[3],
                "human_exit_code": raw[5], "human_classification": detail[3],
                "human_reason": detail[4], "human_log": str(log),
                "human_log_sha256": core.digest(log), "provenance": provenance,
            }
        plan_cases.append({"case": case, "versions": versions, "test": str(test),
            "test_sha256": test_hash, "targets": {TARGET: str(sol)},
            "target_sha256": {TARGET: core.digest(sol)}, "inherited": inherited})
    if expected_keys != set(old_exit) or len(expected_keys) != 1015:
        raise ValueError("Expected exactly the baseline's 1,015 assigned case/version pairs.")
    helpers = {str(p): core.digest(p) for p in sorted(cases_root.glob("*oracle*.py"))}
    return {"model": MODEL, "cases": plan_cases, "helpers": helpers, "timeout_seconds": args.timeout,
        "source_valid": str(args.valid), "source_valid_sha256": core.digest(args.valid),
        "source_baseline": str(args.baseline), "source_baseline_sha256": core.digest(args.baseline),
        "human_policy": "Carried from the recorded baseline, including revised case 790; not rerun.",
        "buggy_policy": "Carried as recorded from baseline; not rerun.",
        "review_policy": "Sol reasons are log-based until question/code/test review. Do not copy Luna sensitivity verdicts.",
        "runner_sha256": core.digest(Path(__file__)),
        "core_sha256": core.digest(Path(core.__file__))}


def collected(plan, batch):
    rows = []
    for entry in plan["cases"]:
        for version in entry["versions"]:
            path = batch / "logs" / entry["case"] / version / f"{TARGET}.json"
            if not path.exists():
                continue
            row = core.read_json(path)
            if (row["case"], row["version"], row["target"]) != (entry["case"], version, TARGET):
                raise ValueError(f"Wrong result identity: {path}")
            if row["test_sha256"] != entry["test_sha256"] or row["target_sha256"] != entry["target_sha256"][TARGET]:
                raise ValueError(f"Stored result has different code/test: {path}")
            if core.digest(Path(row["log"])) != row["log_sha256"]:
                raise ValueError(f"Stored log changed: {path}")
            rows.append(row)
    return rows


def save_results(plan, batch, rows):
    core.write_json(batch / "results.json", {"plan": plan, "records": rows})
    fields = ["case", "version", "target", "raw_result", "exit_code", "classification", "reason",
              "review_required", "seconds", "test", "target_file", "log"]
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: ("'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v)
                         for k, v in row.items()})
    core.atomic_text(batch / "results.csv", "\ufeff" + stream.getvalue())


def export_workbook(plan, batch):
    node = os.environ.get('APR_NODE_EXE') or shutil.which('node')
    if not node:
        raise FileNotFoundError('Node.js missing: set APR_NODE_EXE or add node to PATH. Logs/CSV are saved.')
    exporter = Path(__file__).with_name("export_sol_results.mjs")
    proc = subprocess.run([str(node), str(exporter), str(batch / "results.json")])
    if proc.returncode:
        raise RuntimeError("Workbook export failed. Close Excel and use --report-only; logs/CSV are saved.")
    records = {(r["case"], r["version"]): r for r in collected(plan, batch)}
    book = openpyxl.load_workbook(batch / BOOK_NAME, read_only=True, data_only=True)
    if book.sheetnames != ["Exit results", "Detailed reasons", "Summary"]:
        raise ValueError("Unexpected workbook sheet layout.")
    exits = case_rows(book["Exit results"])
    details = case_rows(book["Detailed reasons"])
    tally = Counter()
    for entry in plan["cases"]:
        for version in entry["versions"]:
            key = (entry["case"], version)
            record = records.get(key, {})
            status = record.get("classification", "NOT TESTED")
            raw = record.get("raw_result", "NOT TESTED")
            old = entry["inherited"][version]
            expected_exit = (entry["case"], version, old["buggy_result"], old["human_raw_result"], raw, old["human_exit_code"], record.get("exit_code"))
            if tuple(exits[key][:7]) != expected_exit:
                raise ValueError(f"Sheet 1 mismatch: {key}")
            if details[key][3:7] != (old["human_classification"], old["human_reason"], status, record.get("reason", "Not yet tested.")):
                raise ValueError(f"Sheet 2 mismatch: {key}")
            if details[key][7] not in (None, ""):
                raise ValueError(f"Unreviewed Sol sensitivity was prefilled: {key}")
            tally[(entry["case"], status)] += 1
    if len(exits) != 1015 or set(exits) != set(details):
        raise ValueError("Workbook case/version coverage mismatch.")
    summary = list(book["Summary"].iter_rows(values_only=True))
    for row in summary[1:74]:
        counts = [tally[(row[0], status)] for status in ("PASS", "ERROR", "FAIL")]
        if list(row[1:5]) != [sum(counts), *counts]:
            raise ValueError(f"Per-case summary mismatch: {row[0]}")
    for sheet in book:
        if any(c.data_type == "e" for row in sheet.iter_rows() for c in row):
            raise ValueError(f"Spreadsheet formula error in {sheet.title}")
    book.close()
    print("Workbook verified: 1,015 assigned pairs; human/buggy baseline preserved; all 73 summaries reconcile.")


def pid_alive(pid):
    if os.name != "nt":
        try:
            os.kill(pid, 0)
            return True
        except ProcessLookupError:
            return False
    from ctypes import wintypes
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
    kernel.OpenProcess.restype = wintypes.HANDLE
    kernel.GetExitCodeProcess.argtypes = [wintypes.HANDLE, ctypes.POINTER(wintypes.DWORD)]
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    handle = kernel.OpenProcess(0x1000, False, pid)
    if not handle:
        return ctypes.get_last_error() != 87
    code = wintypes.DWORD()
    try:
        return not kernel.GetExitCodeProcess(handle, ctypes.byref(code)) or code.value == 259
    finally:
        kernel.CloseHandle(handle)


def prepare_batch(args, plan, batch, snapshots):
    plan_file = batch / "plan.json"
    if plan_file.exists():
        if core.read_json(plan_file) != plan:
            raise RuntimeError("Inputs, baseline, helpers or runner settings changed. Preserve this batch; do not mix runs.")
        if core.read_json(batch / "environments.json") != snapshots:
            raise RuntimeError("An environment changed since this Sol batch began. Do not mix results.")
        return
    if (batch / "environments.json").exists():
        raise RuntimeError("Environment snapshot exists without a plan; inspect this batch.")
    core.write_json(plan_file, plan)
    core.write_json(batch / "environments.json", snapshots)
    old_path = args.root / "gpt-6-luna/testing/first35/environments.json"
    differences = {}
    if old_path.exists():
        old = core.read_json(old_path)
        for version, current in snapshots.items():
            former = old.get(version, {})
            packages = former.get("packages", {})
            changed = {k: [packages.get(k), current["packages"].get(k)]
                       for k in packages.keys() | current["packages"].keys()
                       if packages.get(k) != current["packages"].get(k)}
            if changed or former.get("python") != current.get("python"):
                differences[version] = {"packages": changed, "python": [former.get("python"), current.get("python")]}
    core.write_json(batch / "environment_comparison.json", {
        "historical_human_snapshot": str(old_path), "available": old_path.exists(), "differences": differences})
    if differences:
        print("NOTE: package differences from the historical human run are recorded in environment_comparison.json.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--baseline", type=Path, default=BASELINE)
    parser.add_argument("--valid", type=Path, default=VALID)
    parser.add_argument("--env-root", type=Path, default=Path.home() / "miniconda3/envs")
    parser.add_argument("--timeout", type=int, default=900)
    parser.add_argument("--max-runs", type=int)
    parser.add_argument("--resume-interrupted", action="store_true")
    modes = parser.add_mutually_exclusive_group()
    for option in ("plan", "preflight", "run", "report-only", "status"):
        modes.add_argument("--" + option, action="store_true")
    args = parser.parse_args()
    if args.timeout < 1 or (args.max_runs is not None and args.max_runs < 1):
        parser.error("Timeout and max-runs must be positive.")
    plan = make_plan(args)
    batch = args.root / MODEL / "testing/all73"
    if args.plan or not any((args.preflight, args.run, args.report_only, args.status)):
        for entry in plan["cases"]:
            print(entry["case"] + " | " + ", ".join(entry["versions"]) + " | " + Path(entry["test"]).name)
        print("\n73 cases; 1,015 assigned case/version pairs; 1,015 NEW Sol-only executions.")
        print("Buggy and human results/reasons are carried from:", args.baseline)
        return
    if args.status or args.report_only:
        saved = core.read_json(batch / "plan.json")
        if saved != plan:
            raise RuntimeError("Inputs/settings changed since this batch began.")
        rows = collected(saved, batch)
        print(f"Saved {len(rows)}/1015 Sol executions;", dict(Counter(r["classification"] for r in rows)))
        if args.report_only:
            if (batch / "runner.lock").exists():
                raise RuntimeError("A run lock exists. Finish/resume the run before exporting.")
            save_results(saved, batch, rows)
            export_workbook(saved, batch)
        return
    snapshots = core.probe_environments(plan, args.env_root)
    if args.preflight:
        print("PREFLIGHT OK: source mappings, generated hashes, exact Qiskit versions and pip checks passed. No case tests run.")
        return
    batch.mkdir(parents=True, exist_ok=True)
    lock = batch / "runner.lock"
    if lock.exists():
        info = core.read_json(lock)
        if pid_alive(info["pid"]):
            raise RuntimeError(f"Runner PID {info['pid']} is still active. Do not launch a second run.")
        if not args.resume_interrupted:
            raise RuntimeError("Stale runner lock found. Use --run --resume-interrupted to resume the stopped run.")
        lock.rename(lock.with_name(f"runner.stale_{time.time_ns()}.json"))
    with lock.open("x", encoding="utf-8") as handle:
        json.dump({"pid": os.getpid(), "started_utc": datetime.now(timezone.utc).isoformat()}, handle)
    try:
        prepare_batch(args, plan, batch, snapshots)
        rows = collected(plan, batch)
        done = {(r["case"], r["version"]) for r in rows}
        attempted = 0
        for entry in plan["cases"]:
            for version in entry["versions"]:
                if (entry["case"], version) in done:
                    continue
                if args.max_runs is not None and attempted >= args.max_runs:
                    break
                folder = batch / "logs" / entry["case"] / version
                marker = folder / f"{TARGET}.inflight.json"
                if marker.exists():
                    info = core.read_json(marker)
                    if info.get("pid") and pid_alive(info["pid"]):
                        raise RuntimeError(f"Interrupted child PID {info['pid']} is still active; inspect it before resuming.")
                    if not args.resume_interrupted:
                        raise RuntimeError(f"Interrupted test: {entry['case']}/{version}. Use --resume-interrupted.")
                    stamp = str(time.time_ns())
                    log = folder / f"{TARGET}.txt"
                    if log.exists():
                        log.rename(folder / f"{TARGET}.interrupted_{stamp}.txt")
                    marker.rename(folder / f"{TARGET}.interrupted_{stamp}.json")
                print(f"\n[{len(done)+1}/1015] {entry['case']} / {version} / {TARGET}", flush=True)
                record = core.execute_one(entry, version, TARGET, snapshots[version], batch, args.timeout)
                rows.append(record)
                done.add((entry["case"], version))
                attempted += 1
                save_results(plan, batch, rows)
                print(f"{record['classification']}; exit={record['exit_code']}; {record['seconds']} seconds", flush=True)
            if args.max_runs is not None and attempted >= args.max_runs:
                break
        save_results(plan, batch, rows)
        export_workbook(plan, batch)
        print(f"Saved {len(rows)}/1015 Sol executions to {batch}")
        print(dict(Counter(r["classification"] for r in rows)))
    except KeyboardInterrupt:
        save_results(plan, batch, collected(plan, batch))
        print("\nStopped. Completed results/logs are saved. Resume with --run --resume-interrupted.")
    finally:
        lock.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
